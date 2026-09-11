# MakeWeight — make weight, with the numbers to prove it
# Copyright (C) 2026 Devin Wright (13djwright)
# SPDX-License-Identifier: GPL-3.0-or-later
# This program is free software: you can redistribute it and/or modify it under the terms of the GNU General Public
# License as published by the Free Software Foundation, either version 3 of the License, or (at your option) any
# later version. It is distributed WITHOUT ANY WARRANTY; see the LICENSE file for details.
"""Exports: CSV / xlsx in Devin's sheet layout, print sheet, Bambu presets, 3MF project, robot archive."""
from __future__ import annotations

import csv
import io
import json
import math
import re
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

import numpy as np

from . import meshio, orient, profiles
from .log import log
from .paths import APP_NAME as _APP
from .db import loads, now

COLS = ["Qty", "Price", "Total", "Description", "Purpose / Notes", "Weight (g)", "Total (g)", "Measured (g)", "Source", "Dimensions", "Link"]


def _rows(det: dict):
    """Yield sheet rows in the workbook's shape."""
    yield ["", det["name"], "", "", "", "", "", "", "", "", ""]
    yield ["", "", "", "", "", f"{det['class_name'] or ''} = {det['weight_class_g']:.1f} g", "", "", "", "", ""]
    yield [""] + COLS[:-1] + [COLS[-1]]
    t = det["totals"]
    price_total = sum((it["price"] or 0) * (it["qty"] or 0) for s in det["sections"] for it in s["items"])
    yield ["", "", "", round(price_total, 2), "", "", "", round(t["best_known"], 2), round(t["over_under"], 3), "over(+)/under(-)", ""]
    cfg = next((c["name"] for c in det.get("configs") or [] if c["id"] == det.get("active_config")), None)
    if cfg:
        yield ["", "", "", "", f"Configuration: {cfg}", "lines that are not part of this configuration are listed with a total of 0", "", "", "", "", "", ""]
    for s in det["sections"]:
        first = True
        for it in s["items"]:
            name = s["name"] + ("" if s["counts"] else " (not counted)") if first else ""
            first = False
            purpose = it.get("purpose") or ""
            if it.get("part") and it["part"].get("profile"):
                purpose = (purpose + " · " if purpose else "") + it["part"]["profile"]["string"] + (" · " + it["part"]["filament"]["name"] if it["part"].get("filament") else "")
            yield [name, it["qty"], it.get("price"), round((it["price"] or 0) * (it["qty"] or 0), 2) if it.get("price") else None, it["description"], purpose,
                   round(it["best_grams"], 3) if it["best_grams"] is not None else None, round(it["total_grams"], 3) if it.get("in_total", it["counted"]) else 0,
                   it.get("measured_grams"), ("measured" if it.get("measured_grams") is not None else it.get("est_source") or ""), it.get("dimensions") or "", it.get("link") or ""]
        if first:
            yield [s["name"], "", "", "", "", "", "", "", "", "", "", ""]


def sheet_csv(det: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in _rows(det):
        w.writerow(["" if v is None else v for v in r])
    return buf.getvalue()


def purchase_csv(det: dict) -> str:
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(["Num Required", "Part", "Cost", "Link", "Item Weight (g)", "Gross Weight (g)", "Status"])
    for s in det["sections"]:
        for it in s["items"]:
            if it.get("to_buy") or (it.get("status") in ("planned", "ordered")):
                w.writerow([it["qty"], it["description"], it.get("price") or "", it.get("link") or "", it["best_grams"], it["total_grams"], it.get("status") or ("to buy" if it.get("to_buy") else "")])
    return buf.getvalue()


def sheet_xlsx(app, det: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active; ws.title = "Weight Budget"
    bold = Font(bold=True); head_fill = PatternFill("solid", fgColor="DCE5EE"); sec_fill = PatternFill("solid", fgColor="EEF0EC"); red = Font(color="B6352E", bold=True)
    for r in _rows(det):
        ws.append(["" if v is None else v for v in r])
    ws["B1"].font = Font(bold=True, size=14)
    for c in ws[3]:
        c.font = bold; c.fill = head_fill
    ov = ws.cell(row=4, column=9); ov.font = red if det["totals"]["over_under"] > 0 else Font(color="2E7D4F", bold=True)
    for row in ws.iter_rows(min_row=5):
        if row[0].value:
            row[0].font = bold
            for c in row:
                c.fill = sec_fill
    widths = [16, 6, 8, 9, 38, 44, 11, 10, 12, 10, 16, 40]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A5"
    # printed parts detail
    ps = wb.create_sheet("Printed Parts")
    ps.append(["Part", "Qty", "File", "Filament", "Profile", "Walls", "Top", "Bottom", "Infill %", "Pattern", "Layer", "Orientation", "Slicer g", "Corrected g", "Measured g", "Print time (min)", "Cost $", "Locked", "Role"])
    for c in ps[1]:
        c.font = bold; c.fill = head_fill
    for s in det["sections"]:
        for it in s["items"]:
            p = it.get("part")
            if not p:
                continue
            pr = (p.get("profile") or {}).get("params") or {}
            sl = p.get("slice") or {}
            ps.append([it["description"], it["qty"], (p.get("mesh") or {}).get("filename") or "", (p.get("filament") or {}).get("name") or "", (p.get("profile") or {}).get("string") or "",
                       pr.get("walls"), pr.get("top"), pr.get("bottom"), pr.get("infill"), pr.get("pattern"), pr.get("layer_height"),
                       f"{p['orient'].get('mode', '')} {p['orient'].get('label', '')}".strip(), sl.get("grams"), p.get("corrected_grams"), it.get("measured_grams"),
                       round(sl["print_time_s"] / 60, 1) if sl.get("print_time_s") else None,
                       round(sl["grams"] / 1000 * p["filament"]["cost_per_kg"], 2) if sl.get("grams") and (p.get("filament") or {}).get("cost_per_kg") else None,
                       "yes" if p["locked"] else "", p.get("role")])
    for i, wd in enumerate([30, 5, 36, 20, 30, 6, 5, 7, 8, 10, 6, 18, 9, 11, 11, 7, 10], 1):
        ps.column_dimensions[get_column_letter(i)].width = wd
    # weigh-ins
    wsh = wb.create_sheet("Weigh-ins")
    wsh.append(["Date", "Line", "Grams", "Profile at the time", "Note"])
    for c in wsh[1]:
        c.font = bold; c.fill = head_fill
    for s in det["sections"]:
        for it in s["items"]:
            for w in it.get("weigh_ins", []):
                wsh.append([w.get("date"), it["description"], w["grams"], w.get("profile_string") or "", w.get("note") or ""])
    # events
    ev = wb.create_sheet("Notes")
    ev.append(["Date", "Event", "Placing", "Notes", "Sheet total (g)"])
    for c in ev[1]:
        c.font = bold; c.fill = head_fill
    for e in app.db.q("SELECT * FROM events WHERE robot_id=? ORDER BY date", [det["id"]]):
        ev.append([e["date"], e["title"], e.get("placing") or "", e.get("notes") or "", e.get("total_snapshot_g")])
    ev.column_dimensions["D"].width = 90
    for row in ev.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    # to purchase
    tp = wb.create_sheet("To Purchase")
    for r in csv.reader(io.StringIO(purchase_csv(det))):
        tp.append(r)
    for c in tp[1]:
        c.font = bold; c.fill = head_fill
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def print_sheet_html(app, det: dict) -> str:
    rows = []
    for s in det["sections"]:
        for it in s["items"]:
            p = it.get("part")
            if not p:
                continue
            pr = (p.get("profile") or {}).get("params") or {}
            base = profiles.default_params(pr.get("nozzle", 0.4))
            diffs = []
            for k, label in (("walls", "Wall loops"), ("top", "Top shell layers"), ("bottom", "Bottom shell layers"), ("infill", "Sparse infill density"), ("pattern", "Sparse infill pattern"), ("layer_height", "Layer height")):
                if pr.get(k) != base.get(k):
                    diffs.append(f"{label}: <b>{pr.get(k)}{'%' if k == 'infill' else ''}</b>")
            sl = p.get("slice") or {}
            for md in p.get("modifiers") or []:
                ov = ", ".join(f"{k} {v}{'%' if k == 'infill' else ''}" for k, v in (md.get("params") or {}).items() if v not in (None, ""))
                diffs.append(f"<i>Modifier “{escape(md.get('name') or 'box')}”</i> x {md['min'][0]}…{md['max'][0]}, y {md['min'][1]}…{md['max'][1]}, z {md['min'][2]}…{md['max'][2]} mm: {escape(ov)}")
            rows.append(f"""<tr><td><b>{escape(it['description'])}</b><br><small>{escape((p.get('mesh') or {}).get('filename') or 'no mesh')}</small></td>
              <td>{it['qty']:g}</td><td>{escape((p.get('filament') or {}).get('name') or '')}</td>
              <td>{escape(p['orient'].get('label') or p['orient'].get('mode') or 'auto')}{' · scale ' + str(p['scale']) if p.get('scale') not in (None, 1, 1.0) else ''}{' · mirrored ' + str(p.get('mirror_axis') or 'x').upper() if p.get('mirror') else ''}{' · supports' if ((p.get('print') or {}).get('supports') or {}).get('enabled') else ''}{' · mirrored pair (' + str((p.get('print') or {}).get('pair_mirror')).upper() + ')' if (p.get('print') or {}).get('pair_mirror') else ''}</td>
              <td><code>{escape((p.get('profile') or {}).get('string') or '')}</code><br>{'<br>'.join(diffs) if diffs else '<small>Bambu default</small>'}</td>
              <td class=n>{'' if sl.get('grams') is None else f"{sl['grams']:.1f}"}</td><td class=n>{'' if p.get('corrected_grams') is None else f"{p['corrected_grams']:.1f}"}</td><td class=n>{'' if it.get('measured_grams') is None else f"{it['measured_grams']:.1f}"}</td>
              <td class=n>{'' if not sl.get('print_time_s') else f"{int(sl['print_time_s'] // 3600)}h {int(sl['print_time_s'] % 3600 // 60):02d}m"}</td>
              <td>{'🔒' if p['locked'] else ''}</td></tr>""")
    t = det["totals"]
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>Print sheet · {escape(det['name'])}</title>
<style>body{{font:13px -apple-system,Segoe UI,Roboto,sans-serif;margin:24px;color:#1b1f24}}h1{{margin:0 0 4px}}p{{margin:0 0 14px;color:#4a5159}}table{{border-collapse:collapse;width:100%}}th,td{{border-bottom:1px solid #ccc;padding:6px 8px;text-align:left;vertical-align:top}}th{{font-size:11px;text-transform:uppercase;letter-spacing:.05em;color:#666}}td.n{{text-align:right;font-variant-numeric:tabular-nums}}code{{font-family:ui-monospace,Menlo,monospace;font-size:12px}}small{{color:#7f868e}}@media print{{button{{display:none}}}}</style></head><body>
<button onclick="print()" style="float:right">Print</button>
<h1>{escape(det['name'])} · print sheet</h1>
<p>{escape(det.get('class_name') or '')} · limit {det['weight_class_g']:.1f} g · sheet best-known {t['best_known']:.1f} g ({t['over_under']:+.1f} g) · printed parts {t['printed']:.1f} g · generated {time.strftime('%Y-%m-%d %H:%M')}</p>
<p>In Bambu Studio, start from the <b>0.20mm Standard</b> process for the printer/nozzle and change only the fields listed per part. Supports, brim and skirt are yours to add; weights here exclude them.</p>
<table><thead><tr><th>Part</th><th>Qty</th><th>Filament</th><th>Orientation</th><th>Profile · changes from Bambu default</th><th>Slicer g</th><th>Corrected g</th><th>Measured g</th><th>Print time</th><th></th></tr></thead><tbody>{''.join(rows)}</tbody></table>
</body></html>"""


def bambu_presets_zip(app, det: dict) -> bytes:
    buf = io.BytesIO()
    printer = app.db.get("printers", det.get("printer_id")) if det.get("printer_id") else None
    pname = printer["name"] if printer else "Bambu Lab P1S"
    seen = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for s in det["sections"]:
            for it in s["items"]:
                p = it.get("part")
                if not p or not p.get("profile") or p["profile"]["id"] in seen:
                    continue
                seen.add(p["profile"]["id"])
                prof = p["profile"]
                data = profiles.to_bambu_preset(prof["params"], f"SB {prof['name']}", pname)
                z.writestr(re.sub(r"[^A-Za-z0-9._ -]+", "_", prof["name"]) + ".json", json.dumps(data, indent=2))
        z.writestr("README.txt", "Bambu Studio → (top-left menu) Import → Import Configs… → pick these JSON files.\nThey appear as user process presets named 'SB …'. Assign per object in the object list.\n")
    return buf.getvalue()


def robot_archive(app, rid: int) -> bytes:
    db = app.db
    det = app.robot_detail(rid)
    data = {"format": "makeweight-robot", "version": 1, "exported": time.strftime("%Y-%m-%dT%H:%M:%S"),
            "robot": {k: det[k] for k in ("name", "weight_class_g", "class_name", "margin_g", "nozzle", "notes")} | {"configs": det.get("configs") or [], "active_config": det.get("active_config")},
            "sections": [], "events": db.q("SELECT date,title,placing,notes,total_snapshot_g FROM events WHERE robot_id=?", [rid]),
            "profiles": {}, "filaments": {}, "meshes": {}}
    for s in det["sections"]:
        sec = {"name": s["name"], "counts": s["counts"], "items": []}
        for it in s["items"]:
            item = {k: it.get(k) for k in ("qty", "description", "purpose", "link", "dimensions", "price", "est_grams", "est_source", "status", "needs_reweigh", "to_buy", "counted", "notes", "configs")}
            item["weigh_ins"] = [{k: w.get(k) for k in ("grams", "date", "note", "profile_string")} for w in it.get("weigh_ins", [])]
            p = it.get("part")
            if p:
                part = {k: p.get(k) for k in ("orient", "scale", "role", "locked", "constraints", "mirror", "notes", "modifiers", "print")}
                if p.get("profile"):
                    part["profile"] = p["profile"]["name"]; data["profiles"][p["profile"]["name"]] = p["profile"]["params"]
                if p.get("filament"):
                    f = p["filament"]; part["filament"] = f["name"]; data["filaments"][f["name"]] = {k: f.get(k) for k in ("material", "density", "flow", "color", "cost_per_kg")}
                if p.get("mesh"):
                    m = db.get("meshes", p["mesh"]["id"])
                    part["mesh"] = m["sha256"]; data["meshes"][m["sha256"]] = m["filename"]
                item["part"] = part
            sec["items"].append(item)
        data["sections"].append(sec)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("robot.json", json.dumps(data, indent=1, default=str))
        for sha, fname in data["meshes"].items():
            m = db.one("SELECT path FROM meshes WHERE sha256=?", [sha])
            mp = None
            try:
                mp = meshio.mesh_path(m) if m else None
            except FileNotFoundError:
                mp = None
            if mp:
                z.write(mp, f"meshes/{sha}_{re.sub(r'[^A-Za-z0-9._-]+', '_', fname)}")
    return buf.getvalue()


def import_archive(handler, data: bytes) -> dict:
    """Recreate a robot from a .makeweight.zip archive. Returns the new robot row."""
    app = handler.app; db = app.db
    z = zipfile.ZipFile(io.BytesIO(data))
    meta = json.loads(z.read("robot.json"))
    rb = meta["robot"]
    rid = db.insert("robots", {"name": rb["name"], "weight_class_g": rb["weight_class_g"], "class_name": rb.get("class_name"), "margin_g": rb.get("margin_g") or 4.5,
                               "printer_id": db.setting("default_printer_id"), "nozzle": rb.get("nozzle") or 0.4, "status": "active", "notes": rb.get("notes"), "created": now(), "updated": now(),
                               "configs_json": json.dumps(rb["configs"]) if rb.get("configs") else None, "active_config": rb.get("active_config")})
    fil_ids = {}
    for name, f in meta.get("filaments", {}).items():
        ex = db.one("SELECT id FROM filaments WHERE name=?", [name])
        fil_ids[name] = ex["id"] if ex else db.insert("filaments", {"name": name, "material": f.get("material"), "density": f.get("density") or 1.24, "flow": f.get("flow") or 1.0, "color": f.get("color"), "cost_per_kg": f.get("cost_per_kg"), "correction_json": "{}", "builtin": 0})
    prof_ids = {}
    for name, params in meta.get("profiles", {}).items():
        ph = profiles.profile_hash(profiles.normalize(params))
        ex = next((r for r in db.q("SELECT * FROM profiles") if profiles.profile_hash(profiles.normalize(loads(r["params_json"], {}))) == ph), None)
        prof_ids[name] = ex["id"] if ex else db.insert("profiles", {"name": name, "printer_id": None, "nozzle": params.get("nozzle", 0.4), "params_json": json.dumps(profiles.normalize(params)), "builtin": 0})
    mesh_ids = {}
    for n in z.namelist():
        if n.startswith("meshes/") and not n.endswith("/"):
            sha = n.split("/")[1].split("_")[0]
            res = handler._store_mesh(n.split("_", 1)[1] if "_" in n else "part.stl", z.read(n))
            mesh_ids[sha] = res["meshes"][0]["id"]
    for si, sec in enumerate(meta["sections"]):
        sid = db.insert("sections", {"robot_id": rid, "name": sec["name"], "ord": si, "counts": 1 if sec.get("counts", True) else 0})
        for ii, it in enumerate(sec["items"]):
            iid = db.insert("line_items", {"section_id": sid, "ord": ii, **{k: it.get(k) for k in ("qty", "description", "purpose", "link", "dimensions", "price", "est_grams", "est_source", "status", "notes")},
                                           "needs_reweigh": 1 if it.get("needs_reweigh") else 0, "to_buy": 1 if it.get("to_buy") else 0, "counted": 0 if it.get("counted") is False else 1,
                                           "configs_json": json.dumps(it["configs"]) if it.get("configs") is not None else None})
            for w in it.get("weigh_ins", []):
                db.insert("weigh_ins", {"line_item_id": iid, "grams": w["grams"], "date": w.get("date"), "note": w.get("note"), "profile_string": w.get("profile_string")})
            p = it.get("part")
            if p:
                db.insert("printed_parts", {"robot_id": rid, "line_item_id": iid, "mesh_id": mesh_ids.get(p.get("mesh")), "orient_json": json.dumps(p.get("orient") or {"mode": "auto", "quat": [0, 0, 0, 1]}),
                                            "scale": p.get("scale") or 1.0, "filament_id": fil_ids.get(p.get("filament")) or db.setting("default_filament_id"), "profile_id": prof_ids.get(p.get("profile")) or db.setting("default_profile_id"),
                                            "role": p.get("role") or "structure", "locked": 1 if p.get("locked") else 0, "constraints_json": json.dumps(p.get("constraints") or {}), "mirror": 1 if p.get("mirror") else 0, "notes": p.get("notes"), "modifiers_json": json.dumps(p.get("modifiers") or []), "print_json": json.dumps(p.get("print") or {})})
    for e in meta.get("events", []):
        db.insert("events", {"robot_id": rid, **{k: e.get(k) for k in ("date", "title", "placing", "notes", "total_snapshot_g")}})
    for p in db.q("SELECT id FROM printed_parts WHERE robot_id=?", [rid]):
        app.current_slice_for_part(p["id"])
    return db.get("robots", rid)


# ------------------------------------------------------------------ 3MF
def _object_overrides(params: dict) -> dict:
    """Per-object Bambu keys for one part's profile (what Bambu Studio shows under the object's settings)."""
    if not params:
        return {}
    p = profiles.normalize(params)
    lw = p.get("line_widths") or {}
    out = {
        "wall_loops": str(int(p["walls"])), "top_shell_layers": str(int(p["top"])), "bottom_shell_layers": str(int(p["bottom"])),
        "top_shell_thickness": f"{p['top_min_thickness']:g}", "bottom_shell_thickness": f"{p['bottom_min_thickness']:g}",
        "sparse_infill_density": f"{p['infill']:g}%", "sparse_infill_pattern": profiles.bambu_pattern(p),
        "layer_height": f"{p['layer_height']:g}", "initial_layer_print_height": f"{p['first_layer_height']:g}",
        "infill_wall_overlap": f"{p['infill_wall_overlap']:g}%", "only_one_wall_top": "1" if p.get("one_wall_top") else "0",
        "minimum_sparse_infill_area": f"{p.get('min_sparse_area', 15):g}", "infill_direction": f"{p.get('infill_direction', 45):g}",
        "detect_thin_wall": "1" if p.get("thin_walls") else "0", "filter_out_gap_fill": "0" if p.get("gap_fill", True) else "100",
    }
    if lw:
        out.update({"line_width": f"{lw['outer']:g}", "outer_wall_line_width": f"{lw['outer']:g}", "inner_wall_line_width": f"{lw['inner']:g}",
                    "sparse_infill_line_width": f"{lw['infill']:g}", "internal_solid_infill_line_width": f"{lw['solid']:g}",
                    "top_surface_line_width": f"{lw['top']:g}", "initial_layer_line_width": f"{lw['first']:g}"})
    return out


def bambu_3mf(app, det: dict, machine: str | None = None, nozzle: str | None = None) -> bytes:
    """A Bambu Studio project: one plate per printed line (all copies of that line on its plate), every object carrying
    its own walls / shells / infill / layer settings, modifier regions as modifier parts, the robot's filaments as the
    project's filament list, and a full project_settings.config (printer, default process, filaments) built from the
    same presets the app slices with — so opening it in Bambu Studio needs no re-entering of parameters."""
    from . import bambu_engine
    db = app.db
    jobs = app.jobs
    presets = jobs.presets if jobs.engine == "bambu" else None
    if presets is None:
        res = bambu_engine.resources_dir(bambu_engine.locate(db.setting("bambu_path")))
        presets = bambu_engine.Presets(res) if res else None
    robot = db.get("robots", det["id"]) or {}
    printer = db.get("printers", robot["printer_id"]) if robot.get("printer_id") else db.one("SELECT * FROM printers ORDER BY builtin DESC, id LIMIT 1")
    # the plate grid only lines up in Bambu Studio when the project is laid out for the printer it is opened with — the
    # export dialog lets the user pick; default is the robot's printer
    machine = machine if machine in ("P1S", "H2D") else profiles.machine_key(printer["name"] if printer else None)
    nozzle = f"{float(nozzle or robot.get('nozzle') or 0.4):g}"
    # bed size → plate grid
    bed_w, bed_d = 256.0, 256.0
    if presets:
        try:
            area = presets.machine_for(machine, nozzle).get("printable_area") or []
            xs = [float(a.split("x")[0]) for a in area]; ys = [float(a.split("x")[1]) for a in area]
            bed_w, bed_d = max(xs) - min(xs), max(ys) - min(ys)
        except Exception:  # noqa
            pass
    # collect printed lines
    lines = []          # (item, part, tri_oriented, params, filament row, mesh row)
    fil_rows: list[dict] = []
    for sec in det["sections"]:
        for it in sec["items"]:
            pt = it.get("part")
            if not pt or not pt.get("mesh"):
                continue
            m = db.get("meshes", pt["mesh"]["id"])
            tri = meshio.load_mesh(meshio.mesh_path(m))
            t = orient.apply_orientation(tri, pt["orient"], float(pt.get("scale") or 1.0), orient.mirror_of(pt))
            fil = db.get("filaments", pt["filament_id"]) if pt.get("filament_id") else None
            if fil and all(f["id"] != fil["id"] for f in fil_rows):
                fil_rows.append(fil)
            # a dedicated support-interface material (Bambu Support For PLA …) joins the project's filament list
            sp = (pt.get("print") or {}).get("supports") or {}
            srow = bambu_engine.support_filament_row(sp) if sp.get("enabled") else None
            if srow and all(f["id"] != srow["id"] for f in fil_rows):
                fil_rows.append(srow)
            lines.append((it, pt, t, (pt.get("profile") or {}).get("params") or {}, fil, m))
    if not fil_rows:
        f0 = db.one("SELECT * FROM filaments ORDER BY builtin DESC, id LIMIT 1")
        if f0:
            fil_rows.append(f0)
    fil_index = {f["id"]: i + 1 for i, f in enumerate(fil_rows)}
    n_plates = max(1, len(lines))
    cols = max(1, math.ceil(math.sqrt(n_plates)))
    stride_x, stride_y = bed_w * 1.1, bed_d * 1.1          # Bambu Studio's logical plate gap is 1/10 of the plate

    bs_version = str(bambu_engine.version_of(jobs.slicer_cmd) or "") if jobs.engine == "bambu" and jobs.slicer_cmd else ""
    bs_version = bs_version or "02.08.02.61"
    model = io.StringIO()
    model.write('<?xml version="1.0" encoding="UTF-8"?>\n<model unit="millimeter" xml:lang="en-US" xmlns="http://schemas.microsoft.com/3dmanufacturing/core/2015/02" xmlns:BambuStudio="http://schemas.bambulab.com/package/2021">\n')
    # Bambu Studio only applies Metadata/project_settings.config when the file announces itself as a Bambu project
    model.write(f' <metadata name="Application">BambuStudio-{escape(bs_version)}</metadata>\n <metadata name="BambuStudio:3mfVersion">1</metadata>\n')
    model.write(f' <metadata name="Title">{escape(det.get("name") or "robot")}</metadata>\n <metadata name="Description">Exported by {_APP}</metadata>\n <metadata name="CreationDate">{time.strftime("%Y-%m-%d")}</metadata>\n <resources>\n')
    cfg = io.StringIO()
    cfg.write('<?xml version="1.0" encoding="UTF-8"?>\n<config>\n')
    plates: list[list[tuple[int, int]]] = []      # per plate: [(object_id, identify_id)]
    build_items: list[tuple[int, float, float]] = []
    nid = 1

    def write_mesh(oid: int, name: str, tri):
        verts, faces = meshio._indexed(tri, tol=1e-5)
        model.write(f'  <object id="{oid}" name="{escape(name)}" type="model">\n   <mesh>\n    <vertices>\n')
        for v in verts:
            model.write(f'     <vertex x="{v[0]:.5f}" y="{v[1]:.5f}" z="{v[2]:.5f}"/>\n')
        model.write('    </vertices>\n    <triangles>\n')
        for f in faces:
            model.write(f'     <triangle v1="{f[0]}" v2="{f[1]}" v3="{f[2]}"/>\n')
        model.write('    </triangles>\n   </mesh>\n  </object>\n')

    for plate_i, (it, pt, t, params, fil, m) in enumerate(lines):
        row, col = divmod(plate_i, cols)
        ox, oy = col * stride_x, -row * stride_y
        qty = max(1, int(round(float(it.get("qty") or 1))))
        lo, hi = meshio.bbox(t); size = hi - lo
        # copies side by side, the group centred on the plate; the part frame is x/y-centred, z from the bed
        tc = t - np.array([(lo[0] + hi[0]) / 2, (lo[1] + hi[1]) / 2, lo[2]])
        gap = 8.0
        per_row = max(1, int((bed_w - 20) // (size[0] + gap))) if size[0] + gap < bed_w - 20 else 1
        rows_n = math.ceil(qty / per_row)
        gw = min(qty, per_row) * (size[0] + gap) - gap; gh = rows_n * (size[1] + gap) - gap
        mods = pt.get("modifiers") or []
        overrides = _object_overrides(params)
        sp = (pt.get("print") or {}).get("supports") or {}
        if sp.get("enabled"):
            srow = bambu_engine.support_filament_row(sp)
            overrides.update(bambu_engine.support_keys(sp, fil_index.get(srow["id"]) if srow else None))
        plate_objs = []
        pair_axis = str((pt.get("print") or {}).get("pair_mirror") or "").lower()
        pair_axis = pair_axis if pair_axis in ("x", "y") else None
        tcm = meshio.mirror_placed(tc, pair_axis) if pair_axis and qty > 1 else None
        for k in range(qty):
            r_, c_ = divmod(k, per_row)
            cx = ox + bed_w / 2 - gw / 2 + c_ * (size[0] + gap) + size[0] / 2
            cy = oy + bed_d / 2 - gh / 2 + r_ * (size[1] + gap) + size[1] / 2
            mirrored = tcm is not None and k % 2 == 1          # a left/right pair on one line: every other copy is the mirror image
            name = f"{it['description']}{' #' + str(k + 1) if qty > 1 else ''}{' (mirrored)' if mirrored else ''}"
            parts = [(nid, "normal_part", m.get("filename") or name, tcm if mirrored else tc, {})]
            nid += 1
            for md in mods:
                if md.get("min") and md.get("max"):
                    from .jobs import modifier_settings
                    mn, mx = list(md["min"]), list(md["max"])
                    if mirrored:
                        i_ = "xy".index(pair_axis); mn[i_], mx[i_] = -md["max"][i_], -md["min"][i_]
                    parts.append((nid, "modifier_part", md.get("name") or "modifier", meshio.box_mesh(mn, mx), modifier_settings(md, "bambu")))
                    nid += 1
            for (pid, subtype, pname, ptri, settings) in parts:
                write_mesh(pid, pname, ptri)
            root = nid; nid += 1
            model.write(f'  <object id="{root}" name="{escape(name)}" type="model">\n   <components>\n')
            for (pid, *_r) in parts:
                model.write(f'    <component objectid="{pid}" transform="1 0 0 0 1 0 0 0 1 0 0 0"/>\n')
            model.write('   </components>\n  </object>\n')
            build_items.append((root, cx, cy))
            cfg.write(f'  <object id="{root}">\n    <metadata key="name" value="{escape(name)}"/>\n    <metadata key="extruder" value="{fil_index.get((fil or {}).get("id"), 1)}"/>\n')
            for key, v in overrides.items():
                cfg.write(f'    <metadata key="{key}" value="{escape(str(v))}"/>\n')
            for (pid, subtype, pname, ptri, settings) in parts:
                cfg.write(f'    <part id="{pid}" subtype="{subtype}">\n      <metadata key="name" value="{escape(pname)}"/>\n      <metadata key="matrix" value="1 0 0 0 0 1 0 0 0 0 1 0 0 0 0 1"/>\n')
                if subtype == "normal_part":
                    cfg.write(f'      <metadata key="source_file" value="{escape(m.get("filename") or pname)}"/>\n')
                for sk, sv in settings.items():
                    cfg.write(f'      <metadata key="{escape(str(sk))}" value="{escape(str(sv))}"/>\n')
                cfg.write('    </part>\n')
            cfg.write('  </object>\n')
            plate_objs.append((root, root * 100 + k))
        plates.append(plate_objs)
    model.write(' </resources>\n <build>\n')
    for (oid, cx, cy) in build_items:
        model.write(f'  <item objectid="{oid}" transform="1 0 0 0 1 0 0 0 1 {cx:.4f} {cy:.4f} 0" printable="1"/>\n')
    model.write(' </build>\n</model>\n')
    for i, objs in enumerate(plates or [[]]):
        name = escape(lines[i][0]["description"]) if i < len(lines) else ""
        cfg.write(f'  <plate>\n    <metadata key="plater_id" value="{i + 1}"/>\n    <metadata key="plater_name" value="{name}"/>\n    <metadata key="locked" value="false"/>\n')
        for (oid, ident) in objs:
            cfg.write(f'    <model_instance>\n      <metadata key="object_id" value="{oid}"/>\n      <metadata key="instance_id" value="0"/>\n      <metadata key="identify_id" value="{ident}"/>\n    </model_instance>\n')
        cfg.write('  </plate>\n')
    cfg.write('</config>\n')

    # project settings: the presets the app slices with, in Bambu's project shape (filament keys as per-filament arrays)
    project = None
    if presets:
        try:
            dp_id = robot.get("profile_id") or db.setting("default_profile_id")
            default_prof = db.get("profiles", int(dp_id)) if dp_id else None
            if default_prof is None:
                default_prof = db.one("SELECT * FROM profiles ORDER BY builtin DESC, id LIMIT 1")
            dparams = loads((default_prof or {}).get("params_json"), {}) if default_prof else {}
            mach = presets.machine_for(machine, nozzle)
            proc = presets.process_for(dparams or profiles.default_params(float(nozzle)), machine)
            fils = [presets.filament_for(f, machine, nozzle) for f in fil_rows]
            project = {}
            SKIP = ("name", "from", "inherits", "instantiation", "type", "setting_id", "compatible_printers", "compatible_printers_condition",
                    "description", "include", "filament_id")           # preset bookkeeping, not config — Bambu rejects the file on a type mismatch
            for src in (mach, proc):
                for k, v in src.items():
                    if k in SKIP:
                        continue
                    project[k] = v
            fkeys = set().union(*[set(f.keys()) for f in fils]) if fils else set()
            for k in sorted(fkeys):
                if k in SKIP:
                    continue
                vals = []
                for f in fils:
                    v = f.get(k)
                    vals.append((v[0] if isinstance(v, list) and v else v) if v is not None else (fils[0].get(k)[0] if isinstance(fils[0].get(k), list) and fils[0].get(k) else ""))
                project[k] = [str(x) for x in vals]
            project["filament_settings_id"] = [str(f.get("name") or "filament") for f in fil_rows]
            variant = str((mach.get("printer_extruder_variant") or mach.get("extruder_variant_list") or ["Direct Drive Standard"])[0]).split(",")[0]
            project["filament_self_index"] = [str(i + 1) for i in range(len(fil_rows))]
            project["filament_extruder_variant"] = [str(variant)] * len(fil_rows)
            project["filament_colour"] = [str(f.get("color") or "#8FBC8F") for f in fil_rows]
            project["print_settings_id"] = f"{_APP} {profiles.profile_string(dparams)}" if dparams else f"{_APP} default"
            project["printer_settings_id"] = mach.get("name") or ""
            project["version"] = bs_version
            project["name"] = "project_settings"; project["from"] = "project"
        except Exception as e:  # noqa
            log.warning("3MF export: project_settings.config skipped: %s", e)
            project = None
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\n <Default Extension="model" ContentType="application/vnd.ms-package.3dmanufacturing-3dmodel+xml"/>\n <Default Extension="config" ContentType="text/xml"/>\n</Types>\n')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n <Relationship Target="/3D/3dmodel.model" Id="rel0" Type="http://schemas.microsoft.com/3dmanufacturing/2013/01/3dmodel"/>\n</Relationships>\n')
        z.writestr("3D/3dmodel.model", model.getvalue())
        z.writestr("Metadata/model_settings.config", cfg.getvalue())
        if project:
            z.writestr("Metadata/project_settings.config", json.dumps(project, indent=1, ensure_ascii=False))
        z.writestr("Metadata/filaments.txt", "\n".join(f"Filament {i + 1}: {f.get('name')} ({f.get('density')} g/cm3, flow {f.get('flow')})" for i, f in enumerate(fil_rows)))
    return buf.getvalue()
