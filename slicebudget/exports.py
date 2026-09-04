"""Exports: CSV / xlsx in Devin's sheet layout, print sheet, Bambu presets, 3MF project, robot archive."""
from __future__ import annotations

import csv
import io
import json
import re
import time
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from . import meshio, orient, profiles
from .db import loads, now

COLS = ["Qty", "Price", "Total", "Description", "Purpose / Notes", "Weight (g)", "Total (g)", "Measured (g)", "Source", "Dimensions", "Link"]


def _rows(det: dict):
    """Yield sheet rows in the workbook's shape."""
    yield ["", det["name"], "", "", "", "", "", "", "", "", ""]
    yield ["", "", "", "", "", f"{det['class_name'] or ''} = {det['weight_class_g']:.1f} g", "", "", "", "", ""]
    yield [""] + COLS[:-1] + [COLS[-1]]
    t = det["totals"]
    price_total = sum((it["price"] or 0) * (it["qty"] or 0) for s in det["sections"] for it in s["items"])
    yield ["", "", "", round(price_total, 2), "", "", "", round(t["best_known"], 2), round(t["over_under"], 3), "over(+)/under(-)", ""]
    for s in det["sections"]:
        first = True
        for it in s["items"]:
            name = s["name"] + ("" if s["counts"] else " (not counted)") if first else ""
            first = False
            purpose = it.get("purpose") or ""
            if it.get("part") and it["part"].get("profile"):
                purpose = (purpose + " · " if purpose else "") + it["part"]["profile"]["string"] + (" · " + it["part"]["filament"]["name"] if it["part"].get("filament") else "")
            yield [name, it["qty"], it.get("price"), round((it["price"] or 0) * (it["qty"] or 0), 2) if it.get("price") else None, it["description"], purpose,
                   round(it["best_grams"], 3) if it["best_grams"] is not None else None, round(it["total_grams"], 3) if it["counted"] else 0,
                   it.get("measured_grams"), ("measured" if it.get("measured_grams") is not None else it.get("est_source") or ""), it.get("dimensions") or "", it.get("link") or ""]
        if first:
            yield [s["name"], "", "", "", "", "", "", "", "", "", "", ""]


def sheet_csv(det: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    for r in _rows(det):
        w.writerow(["" if v is None else v for v in r])
    return buf.getvalue()


def purchase_csv(det: dict) -> str:
    buf = io.StringIO(); w = csv.writer(buf)
    w.writerow(["Num Required", "Part", "Cost", "Link", "Item Weight (g)", "Gross Weight (g)", "Status"])
    for s in det["sections"]:
        for it in s["items"]:
            if it.get("to_buy") or (it.get("status") in ("planned", "ordered")):
                w.writerow([it["qty"], it["description"], it.get("price") or "", it.get("link") or "", it["best_grams"], it["total_grams"], it.get("status") or ("to buy" if it.get("to_buy") else "")])
    return buf.getvalue()


def sheet_xlsx(app, det: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active; ws.title = "Weight Budget"
    bold = Font(bold=True); head_fill = PatternFill("solid", fgColor="DCE5EE"); sec_fill = PatternFill("solid", fgColor="EEF0EC"); red = Font(color="B6352E", bold=True)
    for r in _rows(det):
        ws.append(["" if v is None else v for v in r])
    ws["B1"].font = Font(bold=True, size=14)
    for c in ws[3]:
        c.font = bold; c.fill = head_fill
    ov = ws.cell(row=4, column=9); ov.font = red if det["totals"]["over_under"] > 0 else Font(color="2E7D4F", bold=True)
    for row in ws.iter_rows(min_row=5):
        if row[0].value:
            row[0].font = bold
            for c in row:
                c.fill = sec_fill
    widths = [16, 6, 8, 9, 38, 44, 11, 10, 12, 10, 16, 40]
    for i, wd in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A5"
    # printed parts detail
    ps = wb.create_sheet("Printed Parts")
    ps.append(["Part", "Qty", "File", "Filament", "Profile", "Walls", "Top", "Bottom", "Infill %", "Pattern", "Layer", "Orientation", "Slicer g", "Corrected g", "Measured g", "Locked", "Role"])
    for c in ps[1]:
        c.font = bold; c.fill = head_fill
    for s in det["sections"]:
        for it in s["items"]:
            p = it.get("part")
            if not p:
                continue
            pr = (p.get("profile") or {}).get("params") or {}
            sl = p.get("slice") or {}
            ps.append([it["description"], it["qty"], (p.get("mesh") or {}).get("filename") or "", (p.get("filament") or {}).get("name") or "", (p.get("profile") or {}).get("string") or "",
                       pr.get("walls"), pr.get("top"), pr.get("bottom"), pr.get("infill"), pr.get("pattern"), pr.get("layer_height"),
                       f"{p['orient'].get('mode', '')} {p['orient'].get('label', '')}".strip(), sl.get("grams"), p.get("corrected_grams"), it.get("measured_grams"), "yes" if p["locked"] else "", p.get("role")])
    for i, wd in enumerate([30, 5, 36, 20, 30, 6, 5, 7, 8, 10, 6, 18, 9, 11, 11, 7, 10], 1):
        ps.column_dimensions[get_column_letter(i)].width = wd
    # weigh-ins
    wsh = wb.create_sheet("Weigh-ins")
    wsh.append(["Date", "Line", "Grams", "Profile at the time", "Note"])
    for c in wsh[1]:
        c.font = bold; c.fill = head_fill
    for s in det["sections"]:
        for it in s["items"]:
            for w in it.get("weigh_ins", []):
                wsh.append([w.get("date"), it["description"], w["grams"], w.get("profile_string") or "", w.get("note") or ""])
    # events
    ev = wb.create_sheet("Notes")
    ev.append(["Date", "Event", "Placing", "Notes", "Sheet total (g)"])
    for c in ev[1]:
        c.font = bold; c.fill = head_fill
    for e in app.db.q("SELECT * FROM events WHERE robot_id=? ORDER BY date", [det["id"]]):
        ev.append([e["date"], e["title"], e.get("placing") or "", e.get("notes") or "", e.get("total_snapshot_g")])
    ev.column_dimensions["D"].width = 90
    for row in ev.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    # to purchase
    tp = wb.create_sheet("To Purchase")
    for r in csv.reader(io.StringIO(purchase_csv(det))):
        tp.append(r)
    for c in tp[1]:
        c.font = bold; c.fill = head_fill
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


def print_sheet_html(app, det: dict) -> str:
    rows = []
    for s in det["sections"]:
        for it in s["items"]:
            p = it.get("part")
            if not p:
                continue
            pr = (p.get("profile") or {}).get("params") or {}
            base = profiles.default_params(pr.get("nozzle", 0.4))
            diffs = []
            for k, label in (("walls", "Wall loops"), ("top", "Top shell layers"), ("bottom", "Bottom shell layers"), ("infill", "Sparse infill density"), ("pattern", "Sparse infill pattern"), ("layer_height", "Layer height")):
                if pr.get(k) != base.get(k):
                    diffs.append(f"{label}: <b>{pr.get(k)}{'%' if k == 'infill' else ''}</b>")
            sl = p.get("slice") or {}
            rows.append(f"""<tr><td><b>{escape(it['description'])}</b><br><small>{escape((p.get('mesh') or {}).get('filename') or 'no mesh')}</small></td>
              <td>{it['qty']:g}</td><td>{escape((p.get('filament') or {}).get('name') or '')}</td>
              <td>{escape(p['orient'].get('label') or p['orient'].get('mode') or 'auto')}{' · scale ' + str(p['scale']) if p.get('scale') not in (None, 1, 1.0) else ''}{' · mirrored' if p.get('mirror') else ''}</td>
              <td><code>{escape((p.get('profile') or {}).get('string') or '')}</code><br>{'<br>'.join(diffs) if diffs else '<small>Bambu default</small>'}</td>
              <td class=n>{'' if sl.get('grams') is None else f"{sl['grams']:.1f}"}</td><td class=n>{'' if p.get('corrected_grams') is None else f"{p['corrected_grams']:.1f}"}</td><td class=n>{'' if it.get('measured_grams') is None else f"{it['measured_grams']:.1f}"}</td>
              <td>{'🔒' if p['locked'] else ''}</td></tr>""")
    t = det["totals"]
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>Print sheet · {escape(det['name'])}</title>
<style>body{{font:13px -apple-system,Segoe UI,Roboto,sans-serif;margin:24px;color:#1b1f24}}h1{{margin:0 0 4px}}p{{margin:0 0 14px;color:#4a5159}}table{{border-collapse:collapse;width:100%}}th,td{{border-bottom:1px solid #ccc;padding:6px 8px;text-align:left;vertical-align:top}}th{{font-size:11px;text-transform:uppercase;letter-spacing:.05em;color:#666}}td.n{{text-align:right;font-variant-numeric:tabular-nums}}code{{font-family:ui-monospace,Menlo,monospace;font-size:12px}}small{{color:#7f868e}}@media print{{button{{display:none}}}}</style></head><body>
<button onclick="print()" style="float:right">Print</button>
<h1>{escape(det['name'])} · print sheet</h1>
<p>{escape(det.get('class_name') or '')} · limit {det['weight_class_g']:.1f} g · sheet best-known {t['best_known']:.1f} g ({t['over_under']:+.1f} g) · printed parts {t['printed']:.1f} g · generated {time.strftime('%Y-%m-%d %H:%M')}</p>
<p>In Bambu Studio, start from the <b>0.20mm Standard</b> process for the printer/nozzle and change only the fields listed per part. Supports, brim and skirt are yours to add; weights here exclude them.</p>
<table><thead><tr><th>Part</th><th>Qty</th><th>Filament</th><th>Orientation</th><th>Profile · changes from Bambu default</th><th>Slicer g</th><th>Corrected g</th><th>Measured g</th><th></th></tr></thead><tbody>{''.join(rows)}</tbody></table>
</body></html>"""


def bambu_presets_zip(app, det: dict) -> bytes:
    buf = io.BytesIO()
    printer = app.db.get("printers", det.get("printer_id")) if det.get("printer_id") else None
    pname = printer["name"] if printer else "Bambu Lab P1S"
    seen = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for s in det["sections"]:
            for it in s["items"]:
                p = it.get("part")
                if not p or not p.get("profile") or p["profile"]["id"] in seen:
                    continue
                seen.add(p["profile"]["id"])
                prof = p["profile"]
                data = profiles.to_bambu_preset(prof["params"], f"SB {prof['name']}", pname)
                z.writestr(re.sub(r"[^A-Za-z0-9._ -]+", "_", prof["name"]) + ".json", json.dumps(data, indent=2))
        z.writestr("README.txt", "Bambu Studio → (top-left menu) Import → Import Configs… → pick these JSON files.\nThey appear as user process presets named 'SB …'. Assign per object in the object list.\n")
    return buf.getvalue()


def robot_archive(app, rid: int) -> bytes:
    db = app.db
    det = app.robot_detail(rid)
    data = {"format": "slicebudget-robot", "version": 1, "exported": time.strftime("%Y-%m-%dT%H:%M:%S"),
            "robot": {k: det[k] for k in ("name", "weight_class_g", "class_name", "margin_g", "nozzle", "notes")},
            "sections": [], "events": db.q("SELECT date,title,placing,notes,total_snapshot_g FROM events WHERE robot_id=?", [rid]),
            "profiles": {}, "filaments": {}, "meshes": {}}
    for s in det["sections"]:
        sec = {"name": s["name"], "counts": s["counts"], "items": []}
        for it in s["items"]:
            item = {k: it.get(k) for k in ("qty", "description", "purpose", "link", "dimensions", "price", "est_grams", "est_source", "status", "needs_reweigh", "to_buy", "counted", "notes")}
            item["weigh_ins"] = [{k: w.get(k) for k in ("grams", "date", "note", "profile_string")} for w in it.get("weigh_ins", [])]
            p = it.get("part")
            if p:
                part = {k: p.get(k) for k in ("orient", "scale", "role", "locked", "constraints", "mirror", "notes")}
                if p.get("profile"):
                    part["profile"] = p["profile"]["name"]; data["profiles"][p["profile"]["name"]] = p["profile"]["params"]
                if p.get("filament"):
                    f = p["filament"]; part["filament"] = f["name"]; data["filaments"][f["name"]] = {k: f.get(k) for k in ("material", "density", "flow", "color", "cost_per_kg")}
                if p.get("mesh"):
                    m = db.get("meshes", p["mesh"]["id"])
                    part["mesh"] = m["sha256"]; data["meshes"][m["sha256"]] = m["filename"]
                item["part"] = part
            sec["items"].append(item)
        data["sections"].append(sec)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("robot.json", json.dumps(data, indent=1, default=str))
        for sha, fname in data["meshes"].items():
            m = db.one("SELECT path FROM meshes WHERE sha256=?", [sha])
            if m and Path(m["path"]).exists():
                z.write(m["path"], f"meshes/{sha}_{re.sub(r'[^A-Za-z0-9._-]+', '_', fname)}")
    return buf.getvalue()


def import_archive(handler, data: bytes) -> dict:
    """Recreate a robot from a .slicebudget.zip. Returns the new robot row."""
    app = handler.app; db = app.db
    z = zipfile.ZipFile(io.BytesIO(data))
    meta = json.loads(z.read("robot.json"))
    rb = meta["robot"]
    rid = db.insert("robots", {"name": rb["name"], "weight_class_g": rb["weight_class_g"], "class_name": rb.get("class_name"), "margin_g": rb.get("margin_g") or 4.5,
                               "printer_id": db.setting("default_printer_id"), "nozzle": rb.get("nozzle") or 0.4, "status": "active", "notes": rb.get("notes"), "created": now(), "updated": now()})
    fil_ids = {}
    for name, f in meta.get("filaments", {}).items():
        ex = db.one("SELECT id FROM filaments WHERE name=?", [name])
        fil_ids[name] = ex["id"] if ex else db.insert("filaments", {"name": name, "material": f.get("material"), "density": f.get("density") or 1.24, "flow": f.get("flow") or 1.0, "color": f.get("color"), "cost_per_kg": f.get("cost_per_kg"), "correction_json": "{}", "builtin": 0})
    prof_ids = {}
    for name, params in meta.get("profiles", {}).items():
        ph = profiles.profile_hash(profiles.normalize(params))
        ex = next((r for r in db.q("SELECT * FROM profiles") if profiles.profile_hash(profiles.normalize(loads(r["params_json"], {}))) == ph), None)
        prof_ids[name] = ex["id"] if ex else db.insert("profiles", {"name": name, "printer_id": None, "nozzle": params.get("nozzle", 0.4), "params_json": json.dumps(profiles.normalize(params)), "builtin": 0})
    mesh_ids = {}
    for n in z.namelist():
        if n.startswith("meshes/") and not n.endswith("/"):
            sha = n.split("/")[1].split("_")[0]
            res = handler._store_mesh(n.split("_", 1)[1] if "_" in n else "part.stl", z.read(n))
            mesh_ids[sha] = res["meshes"][0]["id"]
    for si, sec in enumerate(meta["sections"]):
        sid = db.insert("sections", {"robot_id": rid, "name": sec["name"], "ord": si, "counts": 1 if sec.get("counts", True) else 0})
        for ii, it in enumerate(sec["items"]):
            iid = db.insert("line_items", {"section_id": sid, "ord": ii, **{k: it.get(k) for k in ("qty", "description", "purpose", "link", "dimensions", "price", "est_grams", "est_source", "status", "notes")},
                                           "needs_reweigh": 1 if it.get("needs_reweigh") else 0, "to_buy": 1 if it.get("to_buy") else 0, "counted": 0 if it.get("counted") is False else 1})
            for w in it.get("weigh_ins", []):
                db.insert("weigh_ins", {"line_item_id": iid, "grams": w["grams"], "date": w.get("date"), "note": w.get("note"), "profile_string": w.get("profile_string")})
            p = it.get("part")
            if p:
                db.insert("printed_parts", {"robot_id": rid, "line_item_id": iid, "mesh_id": mesh_ids.get(p.get("mesh")), "orient_json": json.dumps(p.get("orient") or {"mode": "auto", "quat": [0, 0, 0, 1]}),
                                            "scale": p.get("scale") or 1.0, "filament_id": fil_ids.get(p.get("filament")) or db.setting("default_filament_id"), "profile_id": prof_ids.get(p.get("profile")) or db.setting("default_profile_id"),
                                            "role": p.get("role") or "structure", "locked": 1 if p.get("locked") else 0, "constraints_json": json.dumps(p.get("constraints") or {}), "mirror": 1 if p.get("mirror") else 0, "notes": p.get("notes")})
    for e in meta.get("events", []):
        db.insert("events", {"robot_id": rid, **{k: e.get(k) for k in ("date", "title", "placing", "notes", "total_snapshot_g")}})
    for p in db.q("SELECT id FROM printed_parts WHERE robot_id=?", [rid]):
        app.current_slice_for_part(p["id"])
    return db.get("robots", rid)


# ------------------------------------------------------------------ 3MF
def bambu_3mf(app, det: dict) -> bytes:
    """A 3MF with every printed part oriented and arranged, plus Bambu per-object settings (beta)."""
    db = app.db
    objects = []  # (id, name, tri, params, filament_idx)
    oid = 1
    fil_names = []
    for s in det["sections"]:
        for it in s["items"]:
            p = it.get("part")
            if not p or not p.get("mesh"):
                continue
            m = db.get("meshes", p["mesh"]["id"])
            tri = meshio.load_mesh(Path(m["path"]))
            t = orient.apply_orientation(tri, p["orient"], float(p.get("scale") or 1.0), bool(p.get("mirror")))
            fname = (p.get("filament") or {}).get("name") or "filament"
            if fname not in fil_names:
                fil_names.append(fname)
            for k in range(int(it["qty"] or 1)):
                oid += 1
                objects.append((oid, f"{it['description']}{' #' + str(k + 1) if it['qty'] > 1 else ''}", t, (p.get("profile") or {}).get("params") or {}, fil_names.index(fname) + 1, m["filename"]))
    # arrange in a grid on a 250x250 plate
    placed = []
    x = y = 10.0; row_h = 0.0
    for o in objects:
        lo, hi = meshio.bbox(o[2]); sx, sy = hi[0] - lo[0], hi[1] - lo[1]
        if x + sx > 246 and x > 10:
            x = 10.0; y += row_h + 6; row_h = 0.0
        placed.append((o, x - lo[0], y - lo[1]))
        x += sx + 6; row_h = max(row_h, sy)
    model = io.StringIO()
    model.write('<?xml version="1.0" encoding="UTF-8"?>\n<model unit="millimeter" xml:lang="en-US" xmlns="http://schemas.microsoft.com/3dmanufacturing/core/2015/02" xmlns:BambuStudio="http://schemas.bambulab.com/package/2021">\n')
    model.write(' <metadata name="Application">SliceBudget</metadata>\n <metadata name="BambuStudio:3mfVersion">1</metadata>\n <resources>\n')
    for (o, dx, dy) in placed:
        oid_, name, tri, params, ext, fname = o
        verts, faces = meshio._indexed(tri)
        model.write(f'  <object id="{oid_}" name="{escape(name)}" type="model">\n   <mesh>\n    <vertices>\n')
        for v in verts:
            model.write(f'     <vertex x="{v[0]:.4f}" y="{v[1]:.4f}" z="{v[2]:.4f}"/>\n')
        model.write('    </vertices>\n    <triangles>\n')
        for f in faces:
            model.write(f'     <triangle v1="{f[0]}" v2="{f[1]}" v3="{f[2]}"/>\n')
        model.write('    </triangles>\n   </mesh>\n  </object>\n')
    model.write(' </resources>\n <build>\n')
    for (o, dx, dy) in placed:
        model.write(f'  <item objectid="{o[0]}" transform="1 0 0 0 1 0 0 0 1 {dx:.4f} {dy:.4f} 0" printable="1"/>\n')
    model.write(' </build>\n</model>\n')
    cfg = io.StringIO()
    cfg.write('<?xml version="1.0" encoding="UTF-8"?>\n<config>\n')
    for (o, dx, dy) in placed:
        oid_, name, tri, params, ext, fname = o
        cfg.write(f'  <object id="{oid_}">\n    <metadata key="name" value="{escape(name)}"/>\n    <metadata key="extruder" value="{ext}"/>\n')
        if params:
            keys = {"wall_loops": params.get("walls"), "top_shell_layers": params.get("top"), "bottom_shell_layers": params.get("bottom"),
                    "sparse_infill_density": f"{params.get('infill'):g}%" if params.get("infill") is not None else None, "sparse_infill_pattern": params.get("pattern"),
                    "layer_height": params.get("layer_height"), "top_shell_thickness": params.get("top_min_thickness"), "bottom_shell_thickness": params.get("bottom_min_thickness")}
            for k, v in keys.items():
                if v is not None:
                    cfg.write(f'    <metadata key="{k}" value="{escape(str(v))}"/>\n')
        cfg.write(f'    <part id="{oid_}" subtype="normal_part">\n      <metadata key="name" value="{escape(fname)}"/>\n      <metadata key="matrix" value="1 0 0 0 0 1 0 0 0 0 1 0 0 0 0 1"/>\n      <metadata key="source_file" value="{escape(fname)}"/>\n    </part>\n  </object>\n')
    cfg.write('  <plate>\n    <metadata key="plater_id" value="1"/>\n    <metadata key="plater_name" value=""/>\n    <metadata key="locked" value="false"/>\n')
    for (o, dx, dy) in placed:
        cfg.write(f'    <model_instance>\n      <metadata key="object_id" value="{o[0]}"/>\n      <metadata key="instance_id" value="0"/>\n      <metadata key="identify_id" value="{o[0] * 10}"/>\n    </model_instance>\n')
    cfg.write('  </plate>\n</config>\n')
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?>\n<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">\n <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>\n <Default Extension="model" ContentType="application/vnd.ms-package.3dmanufacturing-3dmodel+xml"/>\n <Default Extension="config" ContentType="text/xml"/>\n</Types>\n')
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">\n <Relationship Target="/3D/3dmodel.model" Id="rel0" Type="http://schemas.microsoft.com/3dmanufacturing/2013/01/3dmodel"/>\n</Relationships>\n')
        z.writestr("3D/3dmodel.model", model.getvalue())
        z.writestr("Metadata/model_settings.config", cfg.getvalue())
        z.writestr("Metadata/filaments.txt", "\n".join(f"Extruder {i + 1}: {n}" for i, n in enumerate(fil_names)))
    return buf.getvalue()
